"""
Service de génération de rapports (Excel / PDF)

Construit des rapports de quantités lisibles (feuille complète + nomenclatures
par type d'ouvrage) à partir des éléments déjà résolus par le parseur
(voir parser_service.py). Aucune dépendance à la couche XSLT retirée.
"""

import base64
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List, Optional
from uuid import UUID

from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session
from weasyprint import HTML

from app.models.database import Element as ElementModel, Model, Storey, Tenant
from app.services.costing_service import costing_service
from app.utils.ifc_utils import (
    HIERARCHY_TYPES,
    AREA_CANDIDATES as _AREA_CANDIDATES,
    VOLUME_CANDIDATES as _VOLUME_CANDIDATES,
    WEIGHT_CANDIDATES as _WEIGHT_CANDIDATES,
    flatten_quantities as _flatten_quantities,
    flatten_properties as _flatten_properties,
    pick_quantity_value as _pick_quantity_value,
)

TEMPLATES_DIR = Path(__file__).resolve().parent.parent / "templates" / "reports"

BRAND_GREEN = "2F6F4F"
HEADER_FILL = PatternFill(start_color=BRAND_GREEN, end_color=BRAND_GREEN, fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ZEBRA_FILL = PatternFill(start_color="F2F5F3", end_color="F2F5F3", fill_type="solid")
TITLE_FONT = Font(bold=True, size=14)
SUBTITLE_FONT = Font(color="595959", size=10)
TOTAL_FILL = PatternFill(start_color="E8EFEA", end_color="E8EFEA", fill_type="solid")
TOTAL_FONT = Font(bold=True)

# Nomenclatures pré-formatées par type IFC : types couverts + quantités à afficher
SCHEDULE_DEFINITIONS: List[Dict[str, Any]] = [
    {
        "key": "walls",
        "title": "Nomenclature — Murs",
        "ifc_types": ["IfcWall", "IfcWallStandardCase", "IfcCurtainWall"],
        "quantity_keys": ["Length", "Width", "Height", "area", "volume"],
    },
    {
        "key": "doors",
        "title": "Nomenclature — Portes",
        "ifc_types": ["IfcDoor"],
        "quantity_keys": ["Width", "Height", "area"],
    },
    {
        "key": "windows",
        "title": "Nomenclature — Fenêtres",
        "ifc_types": ["IfcWindow"],
        "quantity_keys": ["Width", "Height", "area"],
    },
]

_PROPERTY_BONUS_COLUMNS = [
    ("IsExternal", "Extérieur"),
    ("FireRating", "Résistance au feu"),
]

# IfcWall, IfcDoor, etc. sont le vocabulaire standard buildingSMART (IFC) : on
# les garde comme identifiants mais on affiche un libellé français à côté
# (ex. "Mur (IfcWall)"), cohérent avec l'explorateur (voir explorerUtils.ts).
_IFC_TYPE_LABELS: Dict[str, str] = {
    "IfcWall": "Mur",
    "IfcWallStandardCase": "Mur",
    "IfcSlab": "Dalle",
    "IfcDoor": "Porte",
    "IfcWindow": "Fenêtre",
    "IfcBeam": "Poutre",
    "IfcColumn": "Poteau",
    "IfcRoof": "Toiture",
    "IfcStair": "Escalier",
    "IfcStairFlight": "Volée d'escalier",
    "IfcRailing": "Garde-corps",
    "IfcCurtainWall": "Mur-rideau",
    "IfcPlate": "Plaque",
    "IfcMember": "Membrure",
    "IfcCovering": "Revêtement",
    "IfcOpeningElement": "Ouverture",
    "IfcBuildingElementProxy": "Élément générique",
    "IfcFooting": "Fondation",
    "IfcPile": "Pieu",
    "IfcRamp": "Rampe",
    "IfcRampFlight": "Volée de rampe",

    # Renfort structurel
    "IfcReinforcingBar": "Armature (barre)",
    "IfcReinforcingMesh": "Treillis d'armature",
    "IfcTendon": "Câble de précontrainte",
    "IfcTendonAnchor": "Ancrage de précontrainte",
    "IfcMechanicalFastener": "Fixation mécanique",
    "IfcFastener": "Fixation",
    "IfcDiscreteAccessory": "Accessoire",

    # CVC (chauffage, ventilation, climatisation)
    "IfcAirTerminal": "Bouche d'air",
    "IfcAirTerminalBox": "Boîte de bouche d'air",
    "IfcAirToAirHeatRecovery": "Récupérateur de chaleur air-air",
    "IfcBoiler": "Chaudière",
    "IfcBurner": "Brûleur",
    "IfcChiller": "Refroidisseur",
    "IfcCoil": "Batterie d'échange",
    "IfcCompressor": "Compresseur",
    "IfcCondenser": "Condenseur",
    "IfcCoolingTower": "Tour de refroidissement",
    "IfcDamper": "Registre / volet",
    "IfcDuctFitting": "Raccord de gaine",
    "IfcDuctSegment": "Tronçon de gaine",
    "IfcDuctSilencer": "Silencieux de gaine",
    "IfcEvaporativeCooler": "Refroidisseur évaporatif",
    "IfcEvaporator": "Évaporateur",
    "IfcFan": "Ventilateur",
    "IfcFilter": "Filtre",
    "IfcHeatExchanger": "Échangeur de chaleur",
    "IfcHumidifier": "Humidificateur",
    "IfcSpaceHeater": "Radiateur",
    "IfcTank": "Réservoir",
    "IfcTubeBundle": "Faisceau tubulaire",
    "IfcUnitaryControlElement": "Élément de régulation unitaire",
    "IfcUnitaryEquipment": "Équipement unitaire CVC",
    "IfcVibrationIsolator": "Isolateur de vibrations",

    # Plomberie / fluides
    "IfcPipeFitting": "Raccord de tuyauterie",
    "IfcPipeSegment": "Tronçon de tuyauterie",
    "IfcPump": "Pompe",
    "IfcValve": "Vanne",
    "IfcInterceptor": "Séparateur",
    "IfcSanitaryTerminal": "Appareil sanitaire",
    "IfcStackTerminal": "Terminal de colonne",
    "IfcWasteTerminal": "Terminal d'évacuation",
    "IfcFireSuppressionTerminal": "Terminal d'extinction incendie",

    # Réseaux génériques (distribution / régulation)
    "IfcDistributionChamberElement": "Chambre de distribution",
    "IfcDistributionControlElement": "Élément de régulation",
    "IfcDistributionElement": "Élément de distribution",
    "IfcDistributionFlowElement": "Élément de réseau",
    "IfcFlowController": "Régulateur de flux",
    "IfcFlowFitting": "Raccord de réseau",
    "IfcFlowInstrument": "Instrument de mesure",
    "IfcFlowMeter": "Compteur de réseau",
    "IfcFlowMovingDevice": "Dispositif de mise en mouvement",
    "IfcFlowSegment": "Tronçon de réseau",
    "IfcFlowStorageDevice": "Dispositif de stockage (réseau)",
    "IfcFlowTerminal": "Terminal de réseau",
    "IfcFlowTreatmentDevice": "Dispositif de traitement (réseau)",
    "IfcEnergyConversionDevice": "Dispositif de conversion d'énergie",
    "IfcMotorConnection": "Raccordement moteur",
    "IfcJunctionBox": "Boîte de jonction",

    # Électricité
    "IfcActuator": "Actionneur",
    "IfcAlarm": "Alarme",
    "IfcAudioVisualAppliance": "Appareil audiovisuel",
    "IfcCableCarrierFitting": "Raccord de chemin de câbles",
    "IfcCableCarrierSegment": "Tronçon de chemin de câbles",
    "IfcCableFitting": "Raccord de câble",
    "IfcCableSegment": "Tronçon de câble",
    "IfcCommunicationsAppliance": "Appareil de communication",
    "IfcController": "Contrôleur",
    "IfcElectricAppliance": "Appareil électrique",
    "IfcElectricDistributionBoard": "Tableau électrique",
    "IfcElectricFlowStorageDevice": "Dispositif de stockage électrique",
    "IfcElectricGenerator": "Générateur électrique",
    "IfcElectricMotor": "Moteur électrique",
    "IfcElectricTimeControl": "Minuterie électrique",
    "IfcEngine": "Moteur thermique",
    "IfcLamp": "Lampe",
    "IfcLightFixture": "Luminaire",
    "IfcOutlet": "Prise électrique",
    "IfcProtectiveDevice": "Dispositif de protection",
    "IfcProtectiveDeviceTrippingUnit": "Unité de déclenchement",
    "IfcSensor": "Capteur",
    "IfcSolarDevice": "Dispositif solaire",
    "IfcSwitchingDevice": "Appareil de commutation",
    "IfcTransformer": "Transformateur",
    "IfcMedicalDevice": "Appareil médical",

    # Mobilier
    "IfcFurnishingElement": "Élément d'ameublement",
    "IfcFurniture": "Mobilier",
    "IfcSystemFurnitureElement": "Élément de mobilier système",

    # Divers / enveloppe
    "IfcChimney": "Cheminée",
    "IfcCivilElement": "Élément de génie civil",
    "IfcGeographicElement": "Élément géographique",
    "IfcProjectionElement": "Élément de saillie",
    "IfcShadingDevice": "Dispositif d'ombrage",
    "IfcSurfaceFeature": "Élément de surface",
    "IfcTransportElement": "Élément de transport",
    "IfcVirtualElement": "Élément virtuel",
    "IfcVoidingFeature": "Élément d'évidement",

    # Variantes structurelles
    "IfcBeamStandardCase": "Poutre",
    "IfcColumnStandardCase": "Poteau",
    "IfcDoorStandardCase": "Porte",
    "IfcMemberStandardCase": "Membrure",
    "IfcOpeningStandardCase": "Ouverture",
    "IfcPlateStandardCase": "Plaque",
    "IfcSlabElementedCase": "Dalle",
    "IfcSlabStandardCase": "Dalle",
    "IfcWallElementedCase": "Mur",
    "IfcWindowStandardCase": "Fenêtre",
    "IfcCooledBeam": "Poutre froide",
    "IfcBuildingElementPart": "Élément de construction (partie)",
    "IfcElementAssembly": "Assemblage d'éléments",

    "IfcSpace": "Espace",
}


_XLSX_COMPATIBLE_LOGO_TYPES = {"image/png", "image/jpeg"}


def _get_tenant(tenant_id: UUID, db: Session) -> Optional[Tenant]:
    return db.query(Tenant).filter(Tenant.id == tenant_id).first()


def _logo_data_uri(tenant: Optional[Tenant]) -> Optional[str]:
    if not tenant or not tenant.logo_data:
        return None
    encoded = base64.b64encode(tenant.logo_data).decode("ascii")
    return f"data:{tenant.logo_content_type or 'application/octet-stream'};base64,{encoded}"


def _translate_ifc_type(ifc_type: str) -> str:
    label = _IFC_TYPE_LABELS.get(ifc_type)
    return f"{label} ({ifc_type})" if label else ifc_type


@dataclass
class ScheduleColumn:
    key: str
    label: str
    numeric: bool = False
    excel_format: Optional[str] = None


@dataclass
class Schedule:
    title: str
    columns: List[ScheduleColumn]
    rows: List[Dict[str, Any]] = field(default_factory=list)


def _format_property_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Oui" if value else "Non"
    return str(value)


def _element_identity(elem: ElementModel, storey_names: Dict[UUID, str]) -> Dict[str, Any]:
    return {
        "ifc_type": _translate_ifc_type(elem.ifc_type),
        "name": elem.name or "",
        "tag": elem.tag or "",
        "storey": storey_names.get(elem.storey_id, "") if elem.storey_id else "",
    }


def _build_flat_schedule(elements: List[ElementModel], storey_names: Dict[UUID, str]) -> Schedule:
    columns = [
        ScheduleColumn("ifc_type", "Type IFC"),
        ScheduleColumn("name", "Nom"),
        ScheduleColumn("tag", "Repère"),
        ScheduleColumn("storey", "Étage"),
        ScheduleColumn("length", "Longueur (m)", numeric=True, excel_format="#,##0.00"),
        ScheduleColumn("width", "Largeur (m)", numeric=True, excel_format="#,##0.00"),
        ScheduleColumn("height", "Hauteur (m)", numeric=True, excel_format="#,##0.00"),
        ScheduleColumn("area", "Surface (m²)", numeric=True, excel_format="#,##0.00"),
        ScheduleColumn("volume", "Volume (m³)", numeric=True, excel_format="#,##0.00"),
        ScheduleColumn("weight", "Poids (kg)", numeric=True, excel_format="#,##0.00"),
    ]
    rows: List[Dict[str, Any]] = []
    for elem in elements:
        flat_q = _flatten_quantities(elem.quantities)
        row = _element_identity(elem, storey_names)
        row.update({
            "length": _pick_quantity_value(flat_q, ["Length"]),
            "width": _pick_quantity_value(flat_q, ["Width"]),
            "height": _pick_quantity_value(flat_q, ["Height"]),
            "area": _pick_quantity_value(flat_q, _AREA_CANDIDATES),
            "volume": _pick_quantity_value(flat_q, _VOLUME_CANDIDATES),
            "weight": _pick_quantity_value(flat_q, _WEIGHT_CANDIDATES),
        })
        rows.append(row)
    return Schedule(title="Tous les éléments", columns=columns, rows=rows)


def _build_type_schedule(
    definition: Dict[str, Any],
    elements: List[ElementModel],
    storey_names: Dict[UUID, str],
) -> Schedule:
    matched = [e for e in elements if e.ifc_type in definition["ifc_types"]]

    columns = [
        ScheduleColumn("name", "Nom"),
        ScheduleColumn("tag", "Repère"),
        ScheduleColumn("storey", "Étage"),
    ]
    quantity_key_map = {
        "Length": ("length", "Longueur (m)"),
        "Width": ("width", "Largeur (m)"),
        "Height": ("height", "Hauteur (m)"),
        "area": ("area", "Surface (m²)"),
        "volume": ("volume", "Volume (m³)"),
    }
    for qk in definition["quantity_keys"]:
        row_key, label = quantity_key_map[qk]
        columns.append(ScheduleColumn(row_key, label, numeric=True, excel_format="#,##0.00"))

    for prop_key, prop_label in _PROPERTY_BONUS_COLUMNS:
        columns.append(ScheduleColumn(f"prop_{prop_key}", prop_label))

    rows: List[Dict[str, Any]] = []
    for elem in matched:
        flat_q = _flatten_quantities(elem.quantities)
        flat_p = _flatten_properties(elem.properties)
        row = {
            "name": elem.name or "",
            "tag": elem.tag or "",
            "storey": storey_names.get(elem.storey_id, "") if elem.storey_id else "",
        }
        for qk in definition["quantity_keys"]:
            row_key, _ = quantity_key_map[qk]
            if qk == "area":
                row[row_key] = _pick_quantity_value(flat_q, _AREA_CANDIDATES)
            elif qk == "volume":
                row[row_key] = _pick_quantity_value(flat_q, _VOLUME_CANDIDATES)
            else:
                row[row_key] = _pick_quantity_value(flat_q, [qk])
        for prop_key, _ in _PROPERTY_BONUS_COLUMNS:
            row[f"prop_{prop_key}"] = _format_property_value(flat_p.get(prop_key))
        rows.append(row)

    return Schedule(title=definition["title"], columns=columns, rows=rows)


class ReportService:
    """Génère des rapports de quantités (Excel/PDF) pour un modèle donné."""

    def __init__(self):
        self._jinja_env = Environment(
            loader=FileSystemLoader(str(TEMPLATES_DIR)),
            autoescape=select_autoescape(["html"]),
        )

    def _load_elements_and_storeys(
        self, model_id: UUID, tenant_id: UUID, db: Session
    ) -> tuple:
        elements = (
            db.query(ElementModel)
            .filter(
                ElementModel.model_id == model_id,
                ElementModel.tenant_id == tenant_id,
                ~ElementModel.ifc_type.in_(HIERARCHY_TYPES),
            )
            .order_by(ElementModel.ifc_type, ElementModel.name)
            .all()
        )
        storey_ids = {e.storey_id for e in elements if e.storey_id}
        storey_names: Dict[UUID, str] = {}
        if storey_ids:
            for storey in db.query(Storey).filter(Storey.element_id.in_(storey_ids)).all():
                storey_names[storey.element_id] = storey.name or ""
        return elements, storey_names

    def _build_schedules(self, elements: List[ElementModel], storey_names: Dict[UUID, str]) -> List[Schedule]:
        schedules = [_build_flat_schedule(elements, storey_names)]
        for definition in SCHEDULE_DEFINITIONS:
            schedules.append(_build_type_schedule(definition, elements, storey_names))
        return schedules

    def generate_excel(self, model_id: UUID, tenant_id: UUID, db: Session) -> bytes:
        model = db.query(Model).filter(Model.id == model_id, Model.tenant_id == tenant_id).first()
        if not model:
            raise ValueError("Modèle non trouvé")

        elements, storey_names = self._load_elements_and_storeys(model_id, tenant_id, db)
        schedules = self._build_schedules(elements, storey_names)

        workbook = Workbook()
        workbook.remove(workbook.active)

        for schedule in schedules:
            sheet = workbook.create_sheet(title=_safe_sheet_name(schedule.title))
            self._write_schedule_sheet(sheet, schedule)

        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def _write_schedule_sheet(self, sheet: Worksheet, schedule: Schedule) -> None:
        for col_idx, column in enumerate(schedule.columns, start=1):
            cell = sheet.cell(row=1, column=col_idx, value=column.label)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row_idx, row in enumerate(schedule.rows, start=2):
            for col_idx, column in enumerate(schedule.columns, start=1):
                value = row.get(column.key)
                cell = sheet.cell(row=row_idx, column=col_idx, value=value)
                if column.excel_format and value is not None:
                    cell.number_format = column.excel_format
                if row_idx % 2 == 0:
                    cell.fill = ZEBRA_FILL

        sheet.freeze_panes = "A2"
        if schedule.rows:
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(schedule.columns))}{len(schedule.rows) + 1}"

        for col_idx, column in enumerate(schedule.columns, start=1):
            longest = len(column.label)
            for row in schedule.rows:
                value = row.get(column.key)
                if value is not None:
                    longest = max(longest, len(str(value)))
            sheet.column_dimensions[get_column_letter(col_idx)].width = min(max(longest + 3, 10), 40)

    def generate_dpgf(self, model_id: UUID, tenant_id: UUID, db: Session) -> bytes:
        """
        Génère la Décomposition du Prix Global et Forfaitaire (DPGF) d'un
        modèle, groupée par lot (gros œuvre, second œuvre, CVC, électricité...
        voir costing_service.LOT_DEFINITIONS) avec un sous-total par lot — la
        structure attendue d'un DQE/DPGF réel, pas une liste plate.
        """
        model = db.query(Model).filter(Model.id == model_id, Model.tenant_id == tenant_id).first()
        if not model:
            raise ValueError("Modèle non trouvé")

        estimate = costing_service.estimate_model(model_id, tenant_id, db)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = _safe_sheet_name("DPGF")

        tenant = _get_tenant(tenant_id, db)
        if tenant and tenant.logo_data and tenant.logo_content_type in _XLSX_COMPATIBLE_LOGO_TYPES:
            try:
                logo_img = XLImage(BytesIO(tenant.logo_data))
                target_height = 50
                if logo_img.height:
                    scale = target_height / logo_img.height
                    logo_img.width = int(logo_img.width * scale)
                    logo_img.height = target_height
                sheet.add_image(logo_img, "H1")
                sheet.row_dimensions[1].height = 38
            except Exception:
                pass  # logo illisible — ne doit pas bloquer la génération du rapport

        sheet.merge_cells("A1:F1")
        sheet["A1"] = "Décomposition du Prix Global et Forfaitaire (DPGF)"
        sheet["A1"].font = TITLE_FONT

        sheet.merge_cells("A2:F2")
        sheet["A2"] = f"{model.name or 'Modèle sans nom'} — généré le {datetime.utcnow().strftime('%d/%m/%Y')}"
        sheet["A2"].font = SUBTITLE_FONT

        headers = ["N°", "Désignation", "Unité", "Quantité", "Prix unitaire (MAD)", "Prix total (MAD)"]
        header_row = 4
        for col_idx, label in enumerate(headers, start=1):
            cell = sheet.cell(row=header_row, column=col_idx, value=label)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")

        row_idx = header_row + 1
        for lot_idx, lot in enumerate(estimate["lots"], start=1):
            if not lot["line_items"]:
                continue

            sheet.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=6)
            lot_cell = sheet.cell(row=row_idx, column=1, value=lot["label"])
            lot_cell.font = Font(bold=True, color="FFFFFF")
            for col in range(1, 7):
                sheet.cell(row=row_idx, column=col).fill = HEADER_FILL
            row_idx += 1

            for item_idx, li in enumerate(lot["line_items"], start=1):
                values = [
                    f"{lot_idx}.{item_idx}", f"{li['label']} ({li['ifc_type']})", li["unit"],
                    li["quantity"], li["unit_price"], li["total"],
                ]
                for col_idx, value in enumerate(values, start=1):
                    cell = sheet.cell(row=row_idx, column=col_idx, value=value)
                    if col_idx in (4, 5, 6):
                        cell.number_format = "#,##0.00"
                    if row_idx % 2 == 0:
                        cell.fill = ZEBRA_FILL
                row_idx += 1

            sheet.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=5)
            subtotal_label_cell = sheet.cell(row=row_idx, column=1, value=f"Sous-total {lot['label']}")
            subtotal_label_cell.font = TOTAL_FONT
            subtotal_label_cell.fill = TOTAL_FILL
            for col in range(2, 6):
                sheet.cell(row=row_idx, column=col).fill = TOTAL_FILL
            subtotal_cell = sheet.cell(row=row_idx, column=6, value=lot["subtotal"])
            subtotal_cell.font = TOTAL_FONT
            subtotal_cell.fill = TOTAL_FILL
            subtotal_cell.number_format = "#,##0.00"
            row_idx += 1

        row_idx += 1
        sheet.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=5)
        total_label_cell = sheet.cell(row=row_idx, column=1, value="Total général")
        total_label_cell.font = TOTAL_FONT
        total_label_cell.fill = TOTAL_FILL
        for col in range(2, 6):
            sheet.cell(row=row_idx, column=col).fill = TOTAL_FILL
        total_cell = sheet.cell(row=row_idx, column=6, value=estimate["grand_total"])
        total_cell.font = TOTAL_FONT
        total_cell.fill = TOTAL_FILL
        total_cell.number_format = "#,##0.00"
        row_idx += 1

        if estimate["unmatched_types"]:
            row_idx += 1
            sheet.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=6)
            sheet.cell(
                row=row_idx, column=1,
                value="Types sans prix renseigné (non inclus ci-dessus) : " + ", ".join(
                    f"{_translate_ifc_type(u['ifc_type'])} ({u['element_count']})"
                    for u in estimate["unmatched_types"]
                ),
            ).font = SUBTITLE_FONT

        widths = [8, 45, 10, 12, 20, 20]
        for col_idx, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(col_idx)].width = width

        sheet.freeze_panes = f"A{header_row + 1}"

        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def generate_pdf(self, model_id: UUID, tenant_id: UUID, db: Session) -> bytes:
        model = db.query(Model).filter(Model.id == model_id, Model.tenant_id == tenant_id).first()
        if not model:
            raise ValueError("Modèle non trouvé")

        elements, storey_names = self._load_elements_and_storeys(model_id, tenant_id, db)
        schedules = self._build_schedules(elements, storey_names)

        stats = model.statistics or {}
        tenant = _get_tenant(tenant_id, db)
        template = self._jinja_env.get_template("model_report.html")
        html_content = template.render(
            model_name=model.name or "Modèle sans nom",
            project_description=model.description,
            generated_at=datetime.utcnow().strftime("%d/%m/%Y %H:%M"),
            ifc_version=stats.get("ifc_version"),
            total_elements=len(elements),
            total_storeys=stats.get("storeys", len(storey_names)),
            total_spaces=stats.get("spaces", 0),
            warnings=[],
            schedules=schedules,
            logo_data_uri=_logo_data_uri(tenant),
        )

        return HTML(string=html_content, base_url=str(TEMPLATES_DIR)).write_pdf()


def _safe_sheet_name(title: str) -> str:
    # Excel limite les noms de feuille à 31 caractères et interdit []:*?/\
    cleaned = "".join(c for c in title if c not in '[]:*?/\\')
    return cleaned[:31]


report_service = ReportService()
