"""
Points d'extrémité pour le catalogue de prix unitaires (avant-métré chiffré)

Gère le CRUD du catalogue de prix, propre à chaque locataire. Plusieurs prix
peuvent exister pour un même type IFC (ex: deux IfcWall d'épaisseurs
différentes) — voir costing_service.py pour la résolution du prix effectif
d'un élément et elements.py pour l'assignation manuelle par élément.
"""

import re
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, status, Query
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session
from uuid import UUID
from typing import Optional

from app.core.dependencies import get_tenant_id, get_db_session
from app.models.database import PriceCatalogItem
from app.services.report_service import _IFC_TYPE_LABELS

router = APIRouter(prefix="/price-catalog", tags=["price-catalog"])

UNIT_PATTERN = "^(m²|m³|ml|u)$"
VALID_UNITS = {"m²", "m³", "ml", "u"}
HIERARCHY_TYPES = {"IfcProject", "IfcSite", "IfcBuilding", "IfcBuildingStorey"}

# Types utilisables dans le catalogue de prix (tout IFC_TYPE_LABELS sauf les
# conteneurs de hiérarchie) — mêmes 131 types que le sélecteur du formulaire
# manuel (voir explorerUtils.ts IFC_TYPE_GROUPS côté frontend).
IMPORTABLE_TYPES = {t: label for t, label in _IFC_TYPE_LABELS.items() if t not in HIERARCHY_TYPES}

_HEADER_FILL = PatternFill(start_color="2F6F4F", end_color="2F6F4F", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


class PriceCatalogItemCreate(BaseModel):
    ifc_type: str = Field(..., max_length=100)
    label: str = Field(..., max_length=255)
    unit: str = Field(..., pattern=UNIT_PATTERN)
    unit_price: float = Field(..., ge=0)
    notes: Optional[str] = None


class PriceCatalogItemUpdate(BaseModel):
    label: Optional[str] = Field(None, max_length=255)
    unit: Optional[str] = Field(None, pattern=UNIT_PATTERN)
    unit_price: Optional[float] = Field(None, ge=0)
    notes: Optional[str] = None


def _serialize(item: PriceCatalogItem) -> dict:
    return {
        "id": str(item.id),
        "ifc_type": item.ifc_type,
        "label": item.label,
        "unit": item.unit,
        "unit_price": float(item.unit_price),
        "notes": item.notes,
        "created_at": item.created_at.isoformat() if item.created_at else None,
        "updated_at": item.updated_at.isoformat() if item.updated_at else None,
    }


@router.get("")
def list_price_catalog(
    ifc_type: Optional[str] = Query(None),
    tenant_id: UUID = Depends(get_tenant_id),
    db: Session = Depends(get_db_session),
):
    """Liste le catalogue de prix du locataire, optionnellement filtré par type IFC."""
    query = db.query(PriceCatalogItem).filter(PriceCatalogItem.tenant_id == tenant_id)
    if ifc_type:
        query = query.filter(PriceCatalogItem.ifc_type == ifc_type)
    items = query.order_by(PriceCatalogItem.ifc_type, PriceCatalogItem.label).all()
    return [_serialize(item) for item in items]


@router.post("", status_code=status.HTTP_201_CREATED)
def create_price_catalog_item(
    payload: PriceCatalogItemCreate,
    tenant_id: UUID = Depends(get_tenant_id),
    db: Session = Depends(get_db_session),
):
    """Ajoute un prix unitaire (plusieurs entrées possibles par type IFC)."""
    item = PriceCatalogItem(tenant_id=tenant_id, **payload.model_dump())
    db.add(item)
    db.commit()
    db.refresh(item)
    return _serialize(item)


@router.put("/{item_id}")
def update_price_catalog_item(
    item_id: UUID,
    payload: PriceCatalogItemUpdate,
    tenant_id: UUID = Depends(get_tenant_id),
    db: Session = Depends(get_db_session),
):
    """Met à jour un prix unitaire."""
    item = db.query(PriceCatalogItem).filter(
        PriceCatalogItem.id == item_id,
        PriceCatalogItem.tenant_id == tenant_id,
    ).first()
    if not item:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Prix non trouvé")

    for field, value in payload.model_dump(exclude_unset=True).items():
        setattr(item, field, value)

    db.commit()
    db.refresh(item)
    return _serialize(item)


@router.delete("/{item_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_price_catalog_item(
    item_id: UUID,
    tenant_id: UUID = Depends(get_tenant_id),
    db: Session = Depends(get_db_session),
):
    """Supprime un prix unitaire."""
    item = db.query(PriceCatalogItem).filter(
        PriceCatalogItem.id == item_id,
        PriceCatalogItem.tenant_id == tenant_id,
    ).first()
    if not item:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Prix non trouvé")
    db.delete(item)
    db.commit()
    return None


@router.get("/import-template")
def download_import_template():
    """
    Génère un classeur Excel modèle pour l'import en masse : une feuille
    « Import » à remplir (Type IFC / Libellé / Unité / Prix unitaire / Notes,
    avec listes déroulantes) et une feuille « Types disponibles » listant les
    131 types IFC utilisables, pour copier-coller depuis un bordereau de prix
    existant sans avoir à connaître les identifiants IFC par cœur.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Import"

    headers = ["Type IFC", "Libellé", "Unité", "Prix unitaire (MAD)", "Notes"]
    for col_idx, label in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col_idx, value=label)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL

    sheet.append(["IfcWall", "Mur brique 20cm", "m²", 180, "Exemple — à modifier ou supprimer"])

    widths = [22, 32, 10, 18, 30]
    for col_idx, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=col_idx).column_letter].width = width
    sheet.freeze_panes = "A2"

    ref_sheet = workbook.create_sheet("Types disponibles")
    ref_sheet.append(["Type IFC", "Libellé"])
    ref_sheet["A1"].font = _HEADER_FONT
    ref_sheet["A1"].fill = _HEADER_FILL
    ref_sheet["B1"].font = _HEADER_FONT
    ref_sheet["B1"].fill = _HEADER_FILL
    for ifc_type, label in sorted(IMPORTABLE_TYPES.items(), key=lambda kv: (kv[1], kv[0])):
        ref_sheet.append([ifc_type, label])
    ref_sheet.column_dimensions["A"].width = 28
    ref_sheet.column_dimensions["B"].width = 30
    last_row = len(IMPORTABLE_TYPES) + 1

    type_validation = DataValidation(
        type="list", formula1=f"'Types disponibles'!$A$2:$A${last_row}", allow_blank=True
    )
    type_validation.error = "Choisissez un type dans la liste (voir feuille « Types disponibles »)."
    type_validation.errorTitle = "Type IFC invalide"
    sheet.add_data_validation(type_validation)
    type_validation.add("A2:A1000")

    unit_validation = DataValidation(type="list", formula1='"m²,m³,ml,u"', allow_blank=True)
    sheet.add_data_validation(unit_validation)
    unit_validation.add("C2:C1000")

    buffer = BytesIO()
    workbook.save(buffer)

    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="modele_import_prix.xlsx"'},
    )


@router.post("/import")
async def import_price_catalog(
    file: UploadFile = File(...),
    tenant_id: UUID = Depends(get_tenant_id),
    db: Session = Depends(get_db_session),
):
    """
    Importe en masse des prix depuis le classeur Excel généré par
    /price-catalog/import-template. Chaque ligne valide crée un nouveau prix,
    ou met à jour un prix existant si (type IFC, libellé) correspond déjà à
    une entrée du catalogue — permet de corriger un fichier et de le
    réimporter sans créer de doublons. Les lignes invalides sont ignorées et
    rapportées individuellement plutôt que de faire échouer tout l'import.
    """
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Format non supporté. Utilisez le modèle Excel (.xlsx) fourni.",
        )

    content = await file.read()
    try:
        workbook = load_workbook(BytesIO(content), data_only=True)
    except Exception:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Fichier Excel illisible.")

    sheet = workbook["Import"] if "Import" in workbook.sheetnames else workbook.active

    existing = {
        (item.ifc_type, item.label.strip().lower()): item
        for item in db.query(PriceCatalogItem).filter(PriceCatalogItem.tenant_id == tenant_id).all()
    }

    created = 0
    updated = 0
    errors: list[dict] = []

    for row_idx, row in enumerate(sheet.iter_rows(min_row=2, max_col=5, values_only=True), start=2):
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue

        ifc_type_raw, label_raw, unit_raw, price_raw, notes_raw = (list(row) + [None] * 5)[:5]

        if ifc_type_raw is None or label_raw is None or unit_raw is None or price_raw is None:
            errors.append({"row": row_idx, "message": "Champs requis manquants (Type IFC, Libellé, Unité, Prix unitaire)"})
            continue

        # Accepte le code brut ("IfcWall") aussi bien qu'un texte collé depuis la
        # feuille de référence ("Mur (IfcWall)") — on cherche le code n'importe
        # où dans la cellule plutôt que d'exiger qu'elle commence exactement par lui.
        type_match = re.search(r"(Ifc[A-Za-z0-9]+)", str(ifc_type_raw).strip())
        ifc_type = type_match.group(1) if type_match else None
        if not ifc_type or ifc_type not in IMPORTABLE_TYPES:
            errors.append({"row": row_idx, "message": f"Type IFC inconnu : « {ifc_type_raw} »"})
            continue

        # Normalise les espaces insécables qu'Excel insère parfois autour de l'unité.
        unit = str(unit_raw).replace("\xa0", " ").strip()
        if unit not in VALID_UNITS:
            errors.append({"row": row_idx, "message": f"Unité invalide : « {unit_raw} » (attendu : m², m³, ml, u)"})
            continue

        try:
            # Tolère le format numérique français (espace/espace insécable comme
            # séparateur de milliers, virgule comme séparateur décimal), ex:
            # "1 234,56" ou "1\xa0234,56", en plus d'un nombre déjà propre.
            price_str = str(price_raw).replace("\xa0", "").replace(" ", "").replace(",", ".")
            price = float(price_str)
            if price < 0:
                raise ValueError()
        except (TypeError, ValueError):
            errors.append({"row": row_idx, "message": f"Prix unitaire invalide : « {price_raw} »"})
            continue

        label = str(label_raw).strip()
        if not label:
            errors.append({"row": row_idx, "message": "Libellé vide"})
            continue

        notes = str(notes_raw).strip() if notes_raw not in (None, "") else None
        key = (ifc_type, label.lower())

        if key in existing:
            item = existing[key]
            item.unit = unit
            item.unit_price = price
            item.notes = notes
            updated += 1
        else:
            item = PriceCatalogItem(
                tenant_id=tenant_id, ifc_type=ifc_type, label=label, unit=unit, unit_price=price, notes=notes,
            )
            db.add(item)
            existing[key] = item
            created += 1

    db.commit()

    return {
        "created": created,
        "updated": updated,
        "errors": errors,
    }
